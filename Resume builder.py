import openai
from docx import Document
from fpdf import FPDF

# Initialize the OpenAI API client with your API key
openai.api_key = "Enter your API key"
def extract_keywords(job_description):
    prompt = (
        f"Job Description: {job_description}\n\n"
        "Extract the all important keywords from the job description:"
    )

    response = openai.ChatCompletion.create(
        model="gpt-3.5-turbo",  # or another available model
        messages=[
            {"role": "system", "content": "You are an assistant that extracts keywords from job descriptions."},
            {"role": "user", "content": prompt},
        ],
        max_tokens=150,
        n=1,
        stop=None,
        temperature=0.5,
    )

    return response.choices[0].message['content'].strip()

def generate_resume_with_keywords(template, keywords):
    prompt = (
        f"Given the resume template:\n{template}\n\n"
        f"and the keywords: {keywords}\n\n"
        "Generate a professional resume that incorporates the keywords into the relevant sections:"
    )

    response = openai.ChatCompletion.create(
        model="gpt-3.5-turbo",  # or another available model
        messages=[
            {"role": "system", "content": "You are an assistant that generates professional resumes."},
            {"role": "user", "content": prompt},
        ],
        max_tokens=1000,
        n=1,
        stop=None,
        temperature=0.7,
    )

    return response.choices[0].message['content'].strip()

def main():
    job_description = input("Please enter the job description: ")
    keywords = extract_keywords(job_description)
    print("\nExtracted Keywords:\n")
    print(keywords)

    # Update the file path to match the location of your resume template file
    document = Document('resume_template.docx')  # Load the Word document
    resume_template = "\n".join([p.text for p in document.paragraphs])  # Extract text from paragraphs

    updated_resume = generate_resume_with_keywords(resume_template, keywords)
    print("\nUpdated Resume:\n")
    print(updated_resume)

    # Generate PDF directly from the updated resume content
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Times", size=12)

    # Replace the en dash character with a hyphen (-) in the updated_resume
    updated_resume = updated_resume.replace("–", "-")

    # Add heading
    pdf.set_font("Times", size=16)
    pdf.cell(200, 10, "Resume", ln=True, align="C")

    # Add content with reduced line spacing
    pdf.set_font("Times", size=12)
    pdf.multi_cell(0, 7, updated_resume)  # Reduced line spacing to 7

    pdf.output("updated_resume.pdf")

    print("PDF generated successfully.")


if __name__ == "__main__":
    main()
